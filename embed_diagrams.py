# embed_diagrams.py - rebuild from v3, replace 4 diagram slides with PNG images
import sys; sys.path.insert(0, r'C:\ai-agent')
from pptx import Presentation
from pptx.util import Inches, Pt
from ppt_helpers import blank_slide, std_slide_top, body_text, LGRAY
from pptx.enum.text import PP_ALIGN

ASSETS = r"C:\ai-agent\ppt_assets"
prs = Presentation(r"C:\Users\viraj\OneDrive\Desktop\QuickLearner_AI_PPT_v3.pptx")

def replace_slide_with_image(prs, slide_idx, title, pg, img_path, caption):
    """Replace content of slide at index with image + header/footer."""
    slide = prs.slides[slide_idx]
    # Remove all existing shapes
    sp_tree = slide.shapes._spTree
    for sp in list(sp_tree):
        sp_tree.remove(sp)
    # Re-add header + image + caption
    std_slide_top(slide, title, pg)
    slide.shapes.add_picture(img_path, Inches(0.4), Inches(1.98), width=Inches(9.2))
    body_text(slide, caption, Inches(0.4), Inches(6.85), Inches(9.2), Inches(0.45),
              size=10, color=LGRAY, align=PP_ALIGN.CENTER)

# Slide 9 (index 8) = System Architecture
replace_slide_with_image(prs, 8, "System Architecture", 8,
    f"{ASSETS}\\arch.png",
    "3-Tier + AI Service: React.js (Presentation) → FastAPI (Application) → Google Gemini API (AI) + SQLite (Data)")

# Slide 10 (index 9) = DFD Level 0
replace_slide_with_image(prs, 9, "Data Flow Diagram – Level 0  (Context Diagram)", 9,
    f"{ASSETS}\\dfd0.png",
    "Context Diagram: Student provides study content → QuickLearner AI System (Process 0) → Gemini API. All data persisted in SQLite (D1).")

# Slide 13 (index 12) = Use Case
replace_slide_with_image(prs, 12, "Use Case Diagram", 12,
    f"{ASSETS}\\usecase.png",
    "10 use cases. Student (primary actor) uses all features. Gemini API (secondary actor) participates in Chat, Quiz, Flashcards, Summarise, and Highlight Notes.")

# Slide 14 (index 13) = Sequence
replace_slide_with_image(prs, 13, "Sequence Diagram – AI Quiz Generation Flow", 13,
    f"{ASSETS}\\sequence.png",
    "13-message sequence: Student submits topic → Frontend POSTs to FastAPI → Gemini returns MCQ JSON → DB stores result → Student answers → Score updated and displayed.")

out = r"C:\Users\viraj\OneDrive\Desktop\QuickLearner_AI_PPT_Final.pptx"
prs.save(out)
print(f"Saved: {out}  ({len(prs.slides)} slides)")
