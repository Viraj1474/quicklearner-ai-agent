from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONT = "Times New Roman"
FOOTER_LEFT = "SITS, B. E. (Computer) 2019 Course, Project Stage II, 2025-26"
HEADER_TITLE = "QuickLearner AI"

doc = Document(r"C:\Users\viraj\OneDrive\Desktop\final.docx")

# ── helpers ──────────────────────────────────────────────────────────────────
def rule(p, side):
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn('w:pBdr')): pPr.remove(old)
    pBdr = OxmlElement('w:pBdr')
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4')
    el.set(qn('w:space'), '1');    el.set(qn('w:color'), '000000')
    pBdr.append(el); pPr.append(pBdr)

def page_field(p, align, fmt=None):
    p.clear(); p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    r = p.add_run(); r.font.name = FONT; r.font.size = Pt(10)
    instr = f' PAGE \\* {fmt} ' if fmt else ' PAGE '
    fc1 = OxmlElement('w:fldChar'); fc1.set(qn('w:fldCharType'), 'begin')
    it  = OxmlElement('w:instrText'); it.text = instr
    it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'end')
    r._r.extend([fc1, it, fc2])

def clear_hf(hf):
    for p in hf.paragraphs: p.clear()

def restart_numbering(sec, start, fmt):
    sectPr = sec._sectPr
    for old in sectPr.findall(qn('w:pgNumType')): sectPr.remove(old)
    el = OxmlElement('w:pgNumType')
    el.set(qn('w:start'), str(start)); el.set(qn('w:fmt'), fmt)
    sectPr.append(el)

# ── insert a section break (sectPr inside pPr) before paragraph index ────────
def insert_section_break_before(doc, para_idx):
    target_p = doc.paragraphs[para_idx]._p
    new_p = OxmlElement('w:p')
    pPr   = OxmlElement('w:pPr')
    sectPr = OxmlElement('w:sectPr')
    # copy page size/margin from current section
    cur = doc.sections[0]._sectPr
    for tag in [qn('w:pgSz'), qn('w:pgMar')]:
        el = cur.find(tag)
        if el is not None:
            import copy; sectPr.append(copy.deepcopy(el))
    pPr.append(sectPr); new_p.append(pPr)
    target_p.addprevious(new_p)

# ── Insert section break before Chapter 1 (paragraph 95) ─────────────────────
# Para 95 = "Chapter 1" — insert break before it
insert_section_break_before(doc, 95)

# Save & reload so python-docx recognises the new sections
tmp = r"C:\Users\viraj\OneDrive\Desktop\final_tmp.docx"
doc.save(tmp)
doc = Document(tmp)
print("Sections:", len(doc.sections))

# ── Section 0 = Cover + Prelim → roman centered footer, no header ─────────────
s0 = doc.sections[0]
s0.header.is_linked_to_previous = False
s0.footer.is_linked_to_previous = False
clear_hf(s0.header)
clear_hf(s0.footer)
# No header on prelim
s0.header.paragraphs[0].clear()
# Roman page number centered at bottom
page_field(s0.footer.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER, fmt='LowerRoman')
restart_numbering(s0, 1, 'lowerRoman')

# ── Section 1 = Body (Chapter 1 onwards) → SITS footer + title header ─────────
s1 = doc.sections[1]
s1.header.is_linked_to_previous = False
s1.footer.is_linked_to_previous = False

# Header: project title right + bottom rule
clear_hf(s1.header)
hp = s1.header.paragraphs[0]
hp.clear(); hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.space_before = Pt(0)
hp.paragraph_format.space_after  = Pt(4)
rule(hp, 'bottom')
r = hp.add_run(HEADER_TITLE); r.font.name = FONT; r.font.size = Pt(10)

# Footer: SITS left | page# right + top rule
clear_hf(s1.footer)
fp = s1.footer.paragraphs[0]
fp.clear(); fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
fp.paragraph_format.space_before = Pt(0)
fp.paragraph_format.space_after  = Pt(0)
rule(fp, 'top')
# right-align tab stop at ~6 inches from left margin
pPr = fp._p.get_or_add_pPr()
tabs = OxmlElement('w:tabs')
tab  = OxmlElement('w:tab')
tab.set(qn('w:val'), 'right'); tab.set(qn('w:pos'), '8640')
tabs.append(tab); pPr.append(tabs)
r1 = fp.add_run(FOOTER_LEFT); r1.font.name = FONT; r1.font.size = Pt(9)
fp.add_run('\t')
r2 = fp.add_run(); r2.font.name = FONT; r2.font.size = Pt(9)
fc1 = OxmlElement('w:fldChar'); fc1.set(qn('w:fldCharType'), 'begin')
it  = OxmlElement('w:instrText'); it.text = ' PAGE '
it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'end')
r2._r.extend([fc1, it, fc2])
restart_numbering(s1, 1, 'decimal')

# ── Save ──────────────────────────────────────────────────────────────────────
import os
os.remove(tmp)
out = r"C:\Users\viraj\OneDrive\Desktop\final_with_pagenumbers.docx"
doc.save(out)
print("Saved:", out)
