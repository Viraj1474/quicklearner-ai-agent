# ppt_helpers.py  – exact format matching GuardneerConnect reference
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Exact colors from reference ───────────────────────────────────────────────
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BLACK  = RGBColor(0x00, 0x00, 0x00)
BLUE   = RGBColor(0x1F, 0x49, 0x7D)   # dark blue for titles (theme accent)
LGRAY  = RGBColor(0x89, 0x89, 0x89)   # light gray for footer text
YELLOW = RGBColor(0xFF, 0xC0, 0x00)   # accent yellow
DBLUE  = RGBColor(0x26, 0x3F, 0x6B)   # darker blue for highlights
LBLUE  = RGBColor(0xD6, 0xE4, 0xF7)   # light blue fill

W = Inches(10)
H = Inches(7.5)

SINHGAD_LOGO = r"C:\ai-agent\ppt_assets\slide2_img2.jpg"   # Sinhgad logo (left)
SPPU_LOGO    = r"C:\ai-agent\ppt_assets\slide2_img3.jpg"   # SPPU logo (right)

def new_prs():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H
    return prs

def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank

def add_logos(slide):
    """Add Sinhgad logo top-left and SPPU logo top-right — exact positions from reference."""
    slide.shapes.add_picture(SINHGAD_LOGO, Inches(0.087), Inches(0.038),
                              Inches(1.714), Inches(1.101))
    slide.shapes.add_picture(SPPU_LOGO,    Inches(8.233), Inches(0.038),
                              Inches(1.767), Inches(1.2))

def add_header_line(slide):
    """Black horizontal divider line below logos — exact from reference."""
    line = slide.shapes.add_shape(1,  # rectangle as line
        Inches(0), Inches(1.139), Inches(10), Pt(1.5))
    line.fill.solid(); line.fill.fore_color.rgb = BLACK
    line.line.fill.background()

def add_college_text(slide):
    """College name header text — exact from reference."""
    tb = slide.shapes.add_textbox(Inches(1.587), Inches(0.126),
                                   Inches(6.825), Inches(0.808))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Sinhgad Institute of Technology & Science, Pune"
    r.font.name = "Times New Roman"; r.font.size = Pt(16)
    r.font.bold = True; r.font.color.rgb = BLACK
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Department of Computer Engineering"
    r2.font.name = "Times New Roman"; r2.font.size = Pt(14)
    r2.font.color.rgb = BLACK

def add_footer(slide, date="28/6/2026", page_num=None):
    """Date bottom-left, page number bottom-right — exact from reference."""
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(7.0),
                                   Inches(2.233), Inches(0.35))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = date
    r.font.name = "Times New Roman"; r.font.size = Pt(11)
    r.font.color.rgb = LGRAY
    if page_num is not None:
        tb2 = slide.shapes.add_textbox(Inches(9.217), Inches(7.0),
                                        Inches(0.6), Inches(0.35))
        tf2 = tb2.text_frame
        p2 = tf2.paragraphs[0]; p2.alignment = PP_ALIGN.RIGHT
        r2 = p2.add_run(); r2.text = str(page_num)
        r2.font.name = "Times New Roman"; r2.font.size = Pt(12)
        r2.font.color.rgb = BLACK

def slide_title(slide, title_text):
    """Slide title — exact style: top area, bold, dark blue, 34pt."""
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(1.157),
                                   Inches(9.0), Inches(0.75))
    tf = tb.text_frame
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = title_text
    r.font.name = "Times New Roman"; r.font.size = Pt(34)
    r.font.bold = True; r.font.color.rgb = BLUE

def body_text(slide, text, left, top, width, height,
              size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = "Times New Roman"; r.font.size = Pt(size)
    r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    return tb

def bullet_list(slide, items, left, top, width, height, size=16, color=BLACK):
    """Add bulleted list matching reference style."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = f"\u2022  {item}"
        r.font.name = "Times New Roman"; r.font.size = Pt(size)
        r.font.color.rgb = color
    return tb

def std_slide_top(slide, title, page_num, date="28/6/2026"):
    """Full standard slide setup: logos + line + college text + title + footer."""
    add_logos(slide)
    add_header_line(slide)
    add_college_text(slide)
    slide_title(slide, title)
    add_footer(slide, date, page_num)

print("ppt_helpers loaded OK")
